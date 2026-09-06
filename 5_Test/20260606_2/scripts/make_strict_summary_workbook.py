from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RUN_DIR = Path(__file__).resolve().parents[1]
REPORT_DIR = RUN_DIR / "reports"
OUT_PATH = REPORT_DIR / "strict_validation_station_summary_20260606_2.xlsx"


SHEETS = {
    "summary": REPORT_DIR / "strict_metric_summary_2006_2022.csv",
    "good_validation": REPORT_DIR / "strict_good_validation_stations.csv",
    "not_good_validation": REPORT_DIR / "strict_not_good_validation_stations.csv",
    "no_validation": REPORT_DIR / "strict_no_validation_observation_stations.csv",
    "all_metrics": REPORT_DIR / "strict_metrics_by_station_2006_2022.csv",
}


def main() -> None:
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        for sheet, path in SHEETS.items():
            df = pd.read_csv(path, encoding="utf-8-sig")
            df.to_excel(writer, sheet_name=sheet, index=False)

    wb = load_workbook(OUT_PATH)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 32)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        if ws.max_row > 1:
            ws.row_dimensions[1].height = 28
    wb.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
